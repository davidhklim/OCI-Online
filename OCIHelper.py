import os
import sys
import pandas as pd  # Import pandas for Excel file handling
from docx import Document
from docx2pdf import convert  # For converting Word to PDF
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from PIL import Image, ImageTk  # Import Pillow modules
import threading

# Determine base directory based on whether the program is frozen (PyInstaller) or running normally
if getattr(sys, 'frozen', False):
    base_dir = os.path.dirname(sys.executable)
else:
    base_dir = os.path.dirname(os.path.abspath(__file__))

def resource_path(relative_path):
    """Get the absolute path to a resource, works for dev and PyInstaller."""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = base_dir  # Use our determined base_dir
    return os.path.join(base_path, relative_path)

def create_cover_letters(template_path, excel_path, output_folder_base):
    """Generate cover letters (.docx) based on the Word template and Excel data."""
    progress_window = tk.Toplevel(landing_page)
    progress_window.title("Generating Cover Letters")
    progress_window.geometry("400x250")
    progress_window.configure(bg="white")

    progress_label = ttk.Label(progress_window, text="Generating cover letters...", font=("Arial", 12), background="white")
    progress_label.pack(pady=10)

    progress_bar = ttk.Progressbar(progress_window, orient="horizontal", length=300, mode="determinate")
    progress_bar.pack(pady=10)

    current_file_label = ttk.Label(progress_window, text="Current file: None", font=("Arial", 10), background="white")
    current_file_label.pack(pady=10)

    cancel_flag = tk.BooleanVar(value=False)
    def cancel_generation():
        cancel_flag.set(True)
        progress_label.config(text="Cancelling... Please wait.")
        progress_window.update_idletasks()

    cancel_button = ttk.Button(progress_window, text="Cancel", command=cancel_generation)
    cancel_button.pack(pady=10)

    try:
        excel_data = pd.ExcelFile(excel_path)
        sheet_names = excel_data.sheet_names
    except FileNotFoundError:
        messagebox.showerror("Error", f"The Excel file '{excel_path}' was not found.")
        progress_window.destroy()
        return False
    except Exception as e:
        messagebox.showerror("Error", f"Error reading the Excel file: {e}")
        progress_window.destroy()
        return False

    # Count total rows across all sheets (only rows where column AP is TRUE)
    total_files = 0
    for sheet_name in sheet_names:
        try:
            data = excel_data.parse(sheet_name)
            # Only count rows where column 'AP' is TRUE (ignoring case and extra spaces)
            for _, row in data.iterrows():
                if str(row.get("AP", "TRUE")).strip().upper() == "TRUE":
                    total_files += 1
        except Exception as e:
            print(f"Error reading sheet '{sheet_name}': {e}")
            continue
    progress_bar["maximum"] = total_files

    # Get the base name of the selected template (without extension)
    template_file_prefix = os.path.splitext(os.path.basename(template_path))[0]

    for sheet_name in sheet_names:
        try:
            data = excel_data.parse(sheet_name)
        except Exception as e:
            print(f"Error reading sheet '{sheet_name}': {e}")
            continue

        # Create output folder for DOCX files for the current sheet
        output_folder = os.path.join(output_folder_base, f"{sheet_name}_docx")
        os.makedirs(output_folder, exist_ok=True)

        # Map Excel column headers to Word template placeholders
        placeholder_map = {
            "Short Name": "Short_Name",
        }

        for index, row in data.iterrows():
            # Check if the user does NOT want to apply (column AP is FALSE)
            if str(row.get("AP", "TRUE")).strip().upper() != "TRUE":
                continue

            if cancel_flag.get():
                progress_window.destroy()
                messagebox.showinfo("Cancelled", "Cover letter generation was cancelled.")
                return False

            try:
                doc = Document(template_path)
                for paragraph in doc.paragraphs:
                    for key, value in row.items():
                        if pd.notna(value):
                            placeholder = f"«{placeholder_map.get(key, key)}»"
                            for run in paragraph.runs:
                                if placeholder in run.text:
                                    run.text = run.text.replace(placeholder, str(value))
                # Retrieve the "Short Name" and trim any extra spaces
                short_name = row.get("Short Name", f"Document_{index + 1}").strip()
                # Build the output file name as: "[template file base name] (short name).docx"
                output_file_name = f"{template_file_prefix} ({short_name}).docx"
                output_path = os.path.join(output_folder, output_file_name)
                doc.save(output_path)
                print(f"Created: {output_path}")

                current_file_label.config(text=f"Current file: {output_file_name}")
                progress_bar["value"] += 1
                progress_window.update_idletasks()
            except Exception as e:
                print(f"Error generating document for row {index}: {e}")

    progress_window.destroy()
    return True

def convert_all_to_pdf(output_folder_base):
    """Automatically convert all generated DOCX files to PDFs."""
    progress_window = tk.Toplevel(landing_page)
    progress_window.title("Converting to PDF")
    progress_window.geometry("400x250")
    progress_window.configure(bg="white")

    progress_label = ttk.Label(progress_window, text="Converting files to PDF...", font=("Arial", 12), background="white")
    progress_label.pack(pady=10)

    progress_bar = ttk.Progressbar(progress_window, orient="horizontal", length=300, mode="determinate")
    progress_bar.pack(pady=10)

    current_file_label = ttk.Label(progress_window, text="Current file: None", font=("Arial", 10), background="white")
    current_file_label.pack(pady=10)

    cancel_flag = tk.BooleanVar(value=False)
    def cancel_conversion():
        cancel_flag.set(True)
        progress_label.config(text="Cancelling... Please wait.")
        progress_window.update_idletasks()

    cancel_button = ttk.Button(progress_window, text="Cancel", command=cancel_conversion)
    cancel_button.pack(pady=10)

    # Find all DOCX folders (ending with "_docx")
    docx_folders = [os.path.join(output_folder_base, d) for d in os.listdir(output_folder_base)
                    if d.endswith("_docx") and os.path.isdir(os.path.join(output_folder_base, d))]
    total_files = 0
    for folder in docx_folders:
        for f in os.listdir(folder):
            if f.endswith(".docx"):
                total_files += 1
    progress_bar["maximum"] = total_files
    file_count = 0

    for folder in docx_folders:
        pdf_folder = folder.replace("_docx", "_PDF")
        os.makedirs(pdf_folder, exist_ok=True)
        for file in os.listdir(folder):
            if cancel_flag.get():
                progress_window.destroy()
                messagebox.showinfo("Cancelled", "PDF conversion was cancelled.")
                landing_page.destroy()
                os._exit(0)
            if file.endswith(".docx"):
                docx_file = os.path.join(folder, file)
                pdf_file = os.path.join(pdf_folder, os.path.splitext(file)[0] + ".pdf")
                try:
                    current_file_label.config(text=f"Current file: {file}")
                    progress_window.update_idletasks()
                    convert(docx_file, pdf_file)
                    print(f"Converted: {pdf_file}")
                except Exception as e:
                    print(f"Error converting {docx_file}: {e}")
                file_count += 1
                progress_bar["value"] = file_count
                progress_window.update_idletasks()

    progress_window.destroy()
    messagebox.showinfo("Success", "Cover letters (DOCX) and PDFs have been generated successfully!")
    # Once conversion is complete, completely exit the program.
    landing_page.destroy()
    os._exit(0)

def run_generation():
    """Combined process to generate DOCX cover letters and convert them to PDFs."""
    landing_page.withdraw()
    template_path = filedialog.askopenfilename(
        title="Select the Word Template",
        filetypes=[("Word Documents", "*.docx")],
        initialdir=base_dir
    )
    if not template_path:
        messagebox.showerror("Error", "No template selected. Exiting...")
        landing_page.deiconify()
        return

    excel_path = os.path.join(base_dir, "Firm_List.xlsx")
    output_folder_base = base_dir

    success = create_cover_letters(template_path, excel_path, output_folder_base)
    if not success:
        landing_page.deiconify()
        return

    # Start PDF conversion in a separate daemon thread so it won't block exit
    pdf_thread = threading.Thread(target=convert_all_to_pdf, args=(output_folder_base,), daemon=True)
    pdf_thread.start()

def on_closing():
    """Handle the window close event to ensure the program exits completely."""
    if messagebox.askokcancel("Quit", "Do you want to quit?"):
        landing_page.destroy()
        os._exit(0)

if __name__ == "__main__":
    landing_page = tk.Tk()
    landing_page.title("Cover Letter Generator")
    landing_page.geometry("800x850")  # Enlarged landing page
    landing_page.minsize(800, 850)
    landing_page.configure(bg="white")

    # Configure ttk styles
    style = ttk.Style()
    style.configure("TFrame", background="white")
    style.configure("TLabel", background="white", font=("Arial", 12))
    style.configure("TButton", font=("Arial", 12))

    frame = ttk.Frame(landing_page, padding=20)
    frame.pack(fill="both", expand=True)

    header = ttk.Label(
        frame,
        text="Welcome to the Cover Letter Generator!",
        font=("Arial", 20, "bold"),
        background="white",
        anchor="center"
    )
    header.pack(pady=(10, 20))

    # Add preamble: Logo and instructions
    logo_path = resource_path("logo.png")  # Use resource_path to locate the logo
    print(f"Looking for logo at: {logo_path}")  # Debugging: Confirm logo path
    if os.path.exists(logo_path):
        try:
            logo_image = Image.open(logo_path)
            # Resize the logo based on a percentage (120% in this case)
            resize_percentage = 120
            width, height = logo_image.size
            new_width = int(width * (resize_percentage / 100))
            new_height = int(height * (resize_percentage / 100))
            logo_image = logo_image.resize((new_width, new_height), Image.Resampling.LANCZOS)
            logo = ImageTk.PhotoImage(logo_image)
            logo_label = tk.Label(frame, image=logo, bg="white")
            logo_label.image = logo  # Keep a reference to avoid garbage collection
            logo_label.pack(pady=10)
        except Exception as e:
            print(f"Error loading logo: {e}")
    else:
        print("Logo file not found!")

    instructions = ttk.Label(
        frame,
        text=(
            "Copying and pasting 40 addresses and details for every hiring recruiter during the OCI process sucks.\n\n"
            "This tool automates that tedious work by taking your compiled firm information (like addresses, hiring manager details, etc.) "
            "and replacing placeholders in your Word template with firm-specific data. \n\n"
            "For example, the placeholder «Short_Name» will be replaced with 'BLG'.\n\n"
        ),
        wraplength=750,
        justify="left",
        font=("Arial", 12)
    )
    instructions.pack(pady=(10, 5))

    important_label = ttk.Label(
        frame,
        text="Important:",
        font=("Arial", 12, "bold"),
        background="white",
        anchor="w"
    )
    important_label.pack(pady=(5, 0), anchor="w")

    remaining_instructions = ttk.Label(
        frame,
        text=(
            "Make sure to COPY AND PASTE the placeholders directly from the official template into your Word file—do not type them manually.\n\n"
            "Steps:\n"
            "1. Customize your Word template by copying and pasting the placeholder fields exactly as shown.\n"
            "2. In the Firm_List.xlsx file, select the firms you wish to apply to by marking the check boxes.\n"
            "3. Select your Word template when prompted.\n\n"
            "Click 'Next' to begin the process."
        ),
        wraplength=750,
        justify="left",
        font=("Arial", 12)
    )
    remaining_instructions.pack(pady=(5, 20))

    # Single "Next" button to start the entire process
    next_button = ttk.Button(frame, text="Next", command=run_generation)
    next_button.pack(pady=(5, 20))

    landing_page.protocol("WM_DELETE_WINDOW", on_closing)
    landing_page.mainloop()
